import re
import sys
import yaml
import time
from pathlib import Path
from typing import Dict, Optional, Any, List, Tuple
from dataclasses import dataclass
from datetime import datetime

from .logger import get_logger
from .models import (
    DocumentInfo, DocumentMetadata, FileType,
    ArchiveStrategy, ConflictStrategy, DuplicateStrategy,
)

logger = get_logger()

ILLEGAL_CHARS_PATTERN = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
WINDOWS_RESERVED_NAMES = {
    "CON", "PRN", "AUX", "NUL",
    "COM1", "COM2", "COM3", "COM4", "COM5", "COM6", "COM7", "COM8", "COM9",
    "LPT1", "LPT2", "LPT3", "LPT4", "LPT5", "LPT6", "LPT7", "LPT8", "LPT9",
}


@dataclass
class ProcessingRules:
    rename_pattern: str = "{project_code}_{date}_{original_name}"
    date_format: str = "%Y%m%d"
    archive_strategy: ArchiveStrategy = ArchiveStrategy.DATE
    archive_date_format: str = "%Y-%m"
    conflict_strategy: ConflictStrategy = ConflictStrategy.RENAME
    project_code_extract_pattern: Optional[str] = None
    project_code_default: str = "DEFAULT"
    target_dir: Optional[Path] = None
    preserve_original: bool = True
    allowed_file_types: List[FileType] = None
    custom_metadata: Dict[str, str] = None
    enable_deduplication: bool = False
    duplicate_strategy: DuplicateStrategy = DuplicateStrategy.SKIP
    duplicate_area: Optional[Path] = None
    use_fast_hash: bool = False
    enable_ocr: bool = False
    ocr_languages: str = "chi_sim+eng"
    ocr_max_pages: int = 10
    ocr_dpi: int = 300
    enable_persistence: bool = True
    db_path: Optional[Path] = None

    def __post_init__(self) -> None:
        if self.allowed_file_types is None:
            self.allowed_file_types = [
                FileType.PDF, FileType.WORD, FileType.EXCEL, FileType.IMAGE
            ]
        if self.custom_metadata is None:
            self.custom_metadata = {}
        if self.duplicate_area is None:
            self.duplicate_area = Path.cwd() / "_duplicates"


class MetadataExtractor:
    @staticmethod
    def extract_pdf_metadata(file_path: Path, metadata: DocumentMetadata) -> None:
        try:
            import PyPDF2
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                metadata.page_count = len(reader.pages)
                if reader.metadata:
                    meta = reader.metadata
                    metadata.title = meta.get("/Title") or metadata.title
                    metadata.author = meta.get("/Author") or metadata.author
                    if meta.get("/CreationDate"):
                        try:
                            date_str = meta["/CreationDate"].replace("D:", "")[:14]
                            metadata.creation_date = datetime.strptime(
                                date_str, "%Y%m%d%H%M%S"
                            )
                        except (ValueError, TypeError):
                            pass
        except ImportError:
            logger.warning("PyPDF2 未安装，无法提取 PDF 元数据")
        except Exception as e:
            logger.warning(f"提取 PDF 元数据失败 {file_path}: {e}")

    @staticmethod
    def extract_word_metadata(file_path: Path, metadata: DocumentMetadata) -> None:
        try:
            from docx import Document
            doc = Document(str(file_path))
            core_props = doc.core_properties
            metadata.title = core_props.title or metadata.title
            metadata.author = core_props.author or metadata.author
            if core_props.created:
                metadata.creation_date = core_props.created
            if core_props.modified:
                metadata.modification_date = core_props.modified
        except ImportError:
            logger.warning("python-docx 未安装，无法提取 Word 元数据")
        except Exception as e:
            logger.warning(f"提取 Word 元数据失败 {file_path}: {e}")

    @staticmethod
    def extract_excel_metadata(file_path: Path, metadata: DocumentMetadata) -> None:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(filename=str(file_path), read_only=True, data_only=True)
            props = wb.properties
            metadata.title = props.title or metadata.title
            metadata.author = props.creator or metadata.author
            if props.created:
                metadata.creation_date = props.created
            if props.modified:
                metadata.modification_date = props.modified
            wb.close()
        except ImportError:
            logger.warning("openpyxl 未安装，无法提取 Excel 元数据")
        except Exception as e:
            logger.warning(f"提取 Excel 元数据失败 {file_path}: {e}")

    @staticmethod
    def extract_image_metadata(file_path: Path, metadata: DocumentMetadata) -> None:
        try:
            from PIL import Image, ExifTags
            with Image.open(file_path) as img:
                metadata.extra["width"] = str(img.width)
                metadata.extra["height"] = str(img.height)
                metadata.extra["mode"] = img.mode
                if img.format:
                    metadata.extra["format"] = img.format
                if hasattr(img, "_getexif") and img._getexif():
                    exif = {
                        ExifTags.TAGS.get(k, k): v
                        for k, v in img._getexif().items()
                    }
                    if exif.get("DateTimeOriginal"):
                        try:
                            metadata.creation_date = datetime.strptime(
                                exif["DateTimeOriginal"], "%Y:%m:%d %H:%M:%S"
                            )
                        except (ValueError, TypeError):
                            pass
        except ImportError:
            logger.warning("Pillow 未安装，无法提取图片元数据")
        except Exception as e:
            logger.warning(f"提取图片元数据失败 {file_path}: {e}")

    @classmethod
    def extract_full_metadata(cls, doc_info: DocumentInfo) -> None:
        file_path = doc_info.source_path
        metadata = doc_info.metadata

        try:
            if metadata.file_type == FileType.PDF:
                cls.extract_pdf_metadata(file_path, metadata)
            elif metadata.file_type == FileType.WORD:
                cls.extract_word_metadata(file_path, metadata)
            elif metadata.file_type == FileType.EXCEL:
                cls.extract_excel_metadata(file_path, metadata)
            elif metadata.file_type == FileType.IMAGE:
                cls.extract_image_metadata(file_path, metadata)
        except Exception as e:
            logger.warning(f"元数据提取出错 {file_path}: {e}")


class RuleEngine:
    def __init__(self, rules: ProcessingRules) -> None:
        self.rules = rules

    @staticmethod
    def sanitize_filename(filename: str) -> str:
        name = ILLEGAL_CHARS_PATTERN.sub("_", filename)
        name = name.rstrip(" .")
        base = Path(name).stem
        suffix = Path(name).suffix
        if base.upper() in WINDOWS_RESERVED_NAMES:
            base = f"_{base}"
        if len(name) > 255:
            max_base_len = 255 - len(suffix)
            base = base[:max_base_len]
        return f"{base}{suffix}"

    def extract_project_code(self, doc_info: DocumentInfo) -> str:
        if not self.rules.project_code_extract_pattern:
            return self.rules.project_code_default

        try:
            pattern = re.compile(self.rules.project_code_extract_pattern)
            match = pattern.search(doc_info.source_path.name)
            if match:
                return match.group(1) if match.groups() else match.group(0)
        except re.error as e:
            logger.warning(f"项目编号正则表达式错误: {e}")

        path_parts = doc_info.source_path.parts
        for part in reversed(path_parts[:-1]):
            match = re.search(r"(PRJ|PROJ|项目)?[-_]?([A-Za-z0-9]{4,})", part, re.IGNORECASE)
            if match:
                return match.group(2).upper()

        return self.rules.project_code_default

    def generate_target_name(self, doc_info: DocumentInfo) -> str:
        metadata = doc_info.metadata
        original_name = doc_info.source_path.stem
        extension = doc_info.source_path.suffix

        if metadata.creation_date:
            date_str = metadata.creation_date.strftime(self.rules.date_format)
        else:
            date_str = datetime.now().strftime(self.rules.date_format)

        project_code = self.extract_project_code(doc_info)
        metadata.project_code = project_code

        title = metadata.title or original_name
        title = self.sanitize_filename(title)

        template_vars = {
            "project_code": project_code,
            "date": date_str,
            "original_name": self.sanitize_filename(original_name),
            "title": title,
            "file_type": metadata.file_type.value,
            "timestamp": datetime.now().strftime("%H%M%S"),
        }

        for key, value in self.rules.custom_metadata.items():
            template_vars[key] = value

        try:
            new_name = self.rules.rename_pattern.format(**template_vars)
        except KeyError as e:
            logger.warning(f"重命名模板变量缺失 {e}，使用默认模式")
            new_name = f"{project_code}_{date_str}_{original_name}"

        new_name = self.sanitize_filename(new_name)
        return f"{new_name}{extension}"

    def generate_archive_path(self, doc_info: DocumentInfo) -> Optional[Path]:
        if self.rules.archive_strategy == ArchiveStrategy.NONE:
            return None

        if not self.rules.target_dir:
            return None

        base_dir = Path(self.rules.target_dir)
        metadata = doc_info.metadata
        sub_path: Path = base_dir

        if self.rules.archive_strategy == ArchiveStrategy.DATE:
            if metadata.creation_date:
                date_folder = metadata.creation_date.strftime(
                    self.rules.archive_date_format
                )
            else:
                date_folder = datetime.now().strftime(
                    self.rules.archive_date_format
                )
            sub_path = base_dir / date_folder

        elif self.rules.archive_strategy == ArchiveStrategy.PROJECT:
            project_code = metadata.project_code or self.rules.project_code_default
            sub_path = base_dir / project_code

        elif self.rules.archive_strategy == ArchiveStrategy.TYPE:
            sub_path = base_dir / metadata.file_type.value

        return sub_path


class RulesLoader:
    @staticmethod
    def load_from_file(config_path: Path) -> ProcessingRules:
        if not config_path.exists():
            raise FileNotFoundError(f"规则配置文件不存在: {config_path}")

        try:
            with open(config_path, "r", encoding="utf-8") as f:
                config = yaml.safe_load(f)
        except yaml.YAMLError as e:
            raise ValueError(f"规则配置文件格式错误: {e}")

        return RulesLoader._parse_config(config)

    @staticmethod
    def _parse_config(config: Dict[str, Any]) -> ProcessingRules:
        rename = config.get("rename", {})
        archive = config.get("archive", {})
        conflict = config.get("conflict", {})
        extraction = config.get("extraction", {})
        general = config.get("general", {})

        archive_strategy = ArchiveStrategy(
            archive.get("strategy", "date").lower()
        )
        conflict_strategy = ConflictStrategy(
            conflict.get("strategy", "rename").lower()
        )

        allowed_types = []
        for t in general.get("allowed_file_types", ["pdf", "word", "excel", "image"]):
            try:
                allowed_types.append(FileType(t.lower()))
            except ValueError:
                logger.warning(f"忽略未知的文件类型: {t}")

        deduplication = config.get("deduplication", {})
        ocr = config.get("ocr", {})
        persistence = config.get("persistence", {})

        duplicate_strategy = DuplicateStrategy(
            deduplication.get("strategy", "skip").lower()
        )

        return ProcessingRules(
            rename_pattern=rename.get("pattern", "{project_code}_{date}_{original_name}"),
            date_format=rename.get("date_format", "%Y%m%d"),
            archive_strategy=archive_strategy,
            archive_date_format=archive.get("date_format", "%Y-%m"),
            conflict_strategy=conflict_strategy,
            project_code_extract_pattern=extraction.get("project_code_pattern"),
            project_code_default=extraction.get("project_code_default", "DEFAULT"),
            target_dir=Path(archive.get("target_dir")) if archive.get("target_dir") else None,
            preserve_original=general.get("preserve_original", True),
            allowed_file_types=allowed_types,
            custom_metadata=general.get("custom_metadata", {}),
            enable_deduplication=deduplication.get("enabled", False),
            duplicate_strategy=duplicate_strategy,
            duplicate_area=Path(deduplication.get("duplicate_area")) if deduplication.get("duplicate_area") else None,
            use_fast_hash=deduplication.get("use_fast_hash", False),
            enable_ocr=ocr.get("enabled", False),
            ocr_languages=ocr.get("languages", "chi_sim+eng"),
            ocr_max_pages=ocr.get("max_pages", 10),
            ocr_dpi=ocr.get("dpi", 300),
            enable_persistence=persistence.get("enabled", True),
            db_path=Path(persistence.get("db_path")) if persistence.get("db_path") else None,
        )

    @staticmethod
    def load_default() -> ProcessingRules:
        return ProcessingRules()


def retry_operation(
    func,
    max_retries: int = 3,
    delay: float = 1.0,
    backoff: float = 2.0,
) -> Any:
    last_exception = None

    for attempt in range(max_retries + 1):
        try:
            return func()
        except (PermissionError, OSError) as e:
            last_exception = e
            if attempt < max_retries:
                wait_time = delay * (backoff ** attempt)
                logger.warning(
                    f"操作失败 (尝试 {attempt + 1}/{max_retries + 1})，"
                    f"{wait_time:.1f}秒后重试: {e}"
                )
                time.sleep(wait_time)
            else:
                logger.error(f"操作在 {max_retries + 1} 次尝试后仍失败: {e}")

    raise last_exception or RuntimeError("操作失败")
